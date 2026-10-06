"""Perfil estructural por formato. No limpia celdas ni asigna variables físicas.

Loaders: ExcelLoader, CsvLoader, TsvLoader, JsonLoader, TextLoader.
No existe un loader por dataset. El dataset 001 usa el mismo ExcelLoader
que cualquier otro libro.
"""

from __future__ import annotations

import csv
import json
import re
import zipfile
from pathlib import Path
from zipfile import BadZipFile

from openpyxl.utils.exceptions import InvalidFileException

from app.research.profiling import (
    WorkbookError,
    cell_kind,
    profile_workbook,
)

_NAME = re.compile(r"[^a-z0-9]+")
_UNIT = re.compile(r"\(([^)]+)\)\s*$")


class StructureError(Exception):
    def __init__(self, code: str):
        self.code = code
        super().__init__(code)


def normalized_name(raw) -> str | None:
    """Identificador de encabezado. No es un nombre físico canónico."""
    if not isinstance(raw, str):
        return None
    text = _UNIT.sub("", raw.strip().lower()).strip()
    text = _NAME.sub("_", text).strip("_")
    return text or None


def neutralize_spreadsheet_formula(value: str) -> str:
    """Protege una exportación futura. No se aplica al archivo crudo."""
    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def _header(position: int, raw_name, unit, primitive: str, null_count: int, empty: bool, constant: bool) -> dict:
    return {
        "position": position,
        "raw_name": raw_name,
        "normalized_name": normalized_name(raw_name) if raw_name else None,
        "unit": unit,
        "primitive_type": primitive,
        "canonical_variable": None,
        "mapping_status": "pending",
        "null_count": null_count,
        "empty": empty,
        "constant": constant,
    }


def _zip_members(path: Path) -> list[zipfile.ZipInfo]:
    try:
        with zipfile.ZipFile(path) as archive:
            return archive.infolist()
    except (BadZipFile, OSError) as exc:
        raise StructureError("corrupt_workbook") from exc


def _expanded_size(path: Path) -> int:
    return sum(item.file_size for item in _zip_members(path))


class ExcelLoader:
    """Lee xlsx/xlsm sin ejecutar macros, fórmulas ni enlaces externos."""

    def __init__(self, fmt: str):
        self.format = fmt

    def warnings(self, path: Path) -> list[str]:
        found = []
        names = [item.filename for item in _zip_members(path)]
        if any(name.endswith("vbaProject.bin") for name in names):
            found.append("macros_not_executed")
        if any("externalLink" in name for name in names):
            found.append("external_links_not_followed")
        return found

    def profile(self, path: Path, max_bytes: int) -> dict:
        if _expanded_size(path) > max_bytes * 8:
            raise StructureError("file_too_large")
        warnings = self.warnings(path)
        facts = _workbook_facts(path)
        try:
            profile = profile_workbook(path)
        except WorkbookError as exc:
            raise StructureError("corrupt_workbook") from exc
        sheets = []
        for sheet in profile["sheets"]:
            fact = facts["sheets"].get(sheet["name"], {})
            blocks = []
            for index, block in enumerate(sheet["blocks"], start=1):
                headers = []
                for column in block["columns"]:
                    headers.append(
                        _header(
                            column["index"],
                            column.get("name"),
                            column.get("unit"),
                            column.get("inferred_type") or "empty",
                            column.get("null_count") or 0,
                            bool(column.get("empty")),
                            column.get("unique_count") == 1 and (column.get("non_null") or 0) > 0,
                        )
                    )
                header_row = block["start_row"] if block["header_row_count"] else None
                blocks.append(
                    {
                        "block_id": f"B{index:02d}",
                        "start_row": block["start_row"],
                        "end_row": block["end_row"],
                        "start_column": 1,
                        "end_column": block["column_count"],
                        "header_row": header_row,
                        "header_row_count": block["header_row_count"],
                        "headers": headers,
                        "data_row_count": block["data_row_count"],
                        "empty_columns": block["empty_columns"],
                        "duplicate_headers": block["duplicate_headers"],
                    }
                )
            empty_columns = sorted(
                {
                    column
                    for block in sheet["blocks"]
                    for column in block["empty_columns"]
                }
            )
            names = [
                header["raw_name"]
                for block in blocks
                for header in block["headers"]
                if header["raw_name"]
            ]
            sheets.append(
                {
                    "name": sheet["name"],
                    "rows": sheet["row_count"],
                    "columns": sheet["column_count"],
                    "max_row": fact.get("max_row"),
                    "max_column": fact.get("max_column"),
                    "content_rows": sheet["data_row_count"],
                    "content_columns": sheet["column_count"],
                    "empty_row_count": sheet["empty_row_count"],
                    "empty_columns": empty_columns,
                    "block_count": sheet["block_count"],
                    "repeated_headers": bool(sheet["block_count"] > 1 or any(block["duplicate_headers"] for block in blocks)),
                    "header_candidates": sorted({name for name in names}),
                    "blocks": blocks,
                }
            )
        return {
            "warnings": warnings,
            "structure": {
                "workbook": {
                    "sheet_count": len(sheets),
                    "sheet_names": [sheet["name"] for sheet in sheets],
                    "formula_count": facts["formula_count"],
                    "merged_cell_count": len(facts["merged_cells"]),
                    "merged_cells": facts["merged_cells"],
                    "cell_counts": facts["cell_counts"],
                    "macros_executed": False,
                    "formulas_evaluated": False,
                    "sheets": sheets,
                }
            },
        }


def _workbook_facts(path: Path) -> dict:
    from openpyxl import load_workbook

    counts = {"number": 0, "text": 0, "datetime": 0, "boolean": 0, "error": 0}
    try:
        workbook = load_workbook(path, read_only=False, data_only=False, keep_vba=False)
    except (InvalidFileException, BadZipFile, OSError, KeyError, ValueError) as exc:
        raise StructureError("corrupt_workbook") from exc
    formulas = 0
    merged = []
    sheets = {}
    try:
        for worksheet in workbook.worksheets:
            sheets[worksheet.title] = {
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
            }
            for item in worksheet.merged_cells.ranges:
                merged.append({"sheet": worksheet.title, "range": str(item)})
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        formulas += 1
                        continue
                    if cell.data_type == "e":
                        counts["error"] += 1
                        continue
                    kind = cell_kind(cell.value)
                    if kind == "number":
                        counts["number"] += 1
                    elif kind == "text":
                        counts["text"] += 1
                    elif kind == "datetime":
                        counts["datetime"] += 1
                    elif kind == "boolean":
                        counts["boolean"] += 1
    except (InvalidFileException, BadZipFile, OSError, KeyError, ValueError) as exc:
        raise StructureError("corrupt_workbook") from exc
    finally:
        workbook.close()
    return {"formula_count": formulas, "merged_cells": merged, "cell_counts": counts, "sheets": sheets}


def _row_is_header(row: list[str]) -> bool:
    kinds = [_literal_kind(cell) for cell in row]
    texts = sum(kind == "text" for kind in kinds)
    numbers = sum(kind == "number" for kind in kinds)
    return bool(texts) and numbers <= texts


def _literal_kind(value: str) -> str:
    text = value.strip()
    if text == "":
        return "empty"
    if text.lower() in {"true", "false"}:
        return "boolean"
    try:
        float(text)
    except ValueError:
        return "text"
    return "number"


def _encoding(path: Path) -> str:
    with path.open("rb") as handle:
        sample = handle.read(65536)
    encoding = "utf-8-sig" if sample.startswith(b"\xef\xbb\xbf") else "utf-8"
    try:
        sample.decode(encoding)
    except UnicodeError as exc:
        raise StructureError("encoding_unknown") from exc
    return encoding


class _DelimitedLoader:
    def __init__(self, fmt: str, preferred: str | None):
        self.format = fmt
        self.preferred = preferred

    def profile(self, path: Path, max_bytes: int) -> dict:
        if path.stat().st_size > max_bytes:
            raise StructureError("file_too_large")
        encoding = _encoding(path)
        with path.open("rb") as handle:
            sample = handle.read(65536)
        text = sample.decode(encoding)
        warnings = []
        delimiter, quote = _sniff(text, self.preferred)
        if self.preferred and delimiter != self.preferred:
            warnings.append("delimiter_differs_from_extension")
        headers, rows_meta = _stream_table(path, encoding, delimiter, quote)
        return {
            "warnings": warnings,
            "structure": {
                "tabular": {
                    "encoding": encoding,
                    "delimiter": delimiter,
                    "quote": quote,
                    "header_row": 1 if headers else None,
                    "physical_row_count": rows_meta["physical_row_count"],
                    "data_row_count": rows_meta["data_row_count"],
                    "column_count": rows_meta["column_count"],
                    "empty_row_count": rows_meta["empty_row_count"],
                    "empty_columns": rows_meta["empty_columns"],
                    "headers": headers,
                }
            },
        }


def _sniff(sample: str, preferred: str | None) -> tuple[str, str]:
    try:
        dialect = csv.Sniffer().sniff(sample or ",", delimiters=",\t;|")
        delimiter = dialect.delimiter
        quote = dialect.quotechar or '"'
    except csv.Error:
        delimiter = preferred or ","
        quote = '"'
    return delimiter, quote


def _stream_table(path: Path, encoding: str, delimiter: str, quote: str) -> tuple[list[dict], dict]:
    physical = 0
    empty_rows = 0
    data_rows = 0
    width = 0
    header_cells = None
    kinds: list[set[str]] = []
    nulls: list[int] = []
    uniques: list[set[str]] = []
    try:
        with path.open("r", encoding=encoding, newline="") as handle:
            reader = csv.reader(handle, delimiter=delimiter, quotechar=quote)
            for physical, row in enumerate(reader, start=1):
                if physical == 1 and _row_is_header(row):
                    header_cells = list(row)
                    width = len(header_cells)
                    kinds = [set() for _ in header_cells]
                    nulls = [0 for _ in header_cells]
                    uniques = [set() for _ in header_cells]
                    continue
                    width = len(row)
                    kinds = [set() for _ in row]
                    nulls = [0 for _ in row]
                    uniques = [set() for _ in row]
                if width == 0:
                    width = len(row)
                    kinds = [set() for _ in row]
                    nulls = [0 for _ in row]
                    uniques = [set() for _ in row]
                if len(row) > width:
                    extra = len(row) - width
                    kinds.extend(set() for _ in range(extra))
                    nulls.extend(0 for _ in range(extra))
                    uniques.extend(set() for _ in range(extra))
                    if header_cells is not None:
                        header_cells.extend("" for _ in range(extra))
                    width = len(row)
                values = list(row) + [""] * (width - len(row))
                if all(not cell.strip() for cell in values):
                    empty_rows += 1
                    continue
                data_rows += 1
                for index, cell in enumerate(values):
                    kind = _literal_kind(cell)
                    if kind == "empty":
                        nulls[index] += 1
                        continue
                    kinds[index].add(kind)
                    if len(uniques[index]) < 3:
                        uniques[index].add(cell.strip())
    except UnicodeError as exc:
        raise StructureError("encoding_unknown") from exc
    headers = []
    empty_columns = []
    for index in range(width):
        present = kinds[index]
        if not present:
            primitive = "empty"
            empty_columns.append(index + 1)
        elif len(present) == 1:
            primitive = next(iter(present))
        else:
            primitive = "mixed"
        raw = None
        if header_cells and index < len(header_cells) and header_cells[index]:
            raw = header_cells[index]
        unit = None
        if raw:
            match = _UNIT.search(raw.strip())
            unit = match.group(1) if match else None
        constant = len(uniques[index]) == 1 and primitive != "empty"
        headers.append(_header(index + 1, raw, unit, primitive, nulls[index] if index < len(nulls) else 0, primitive == "empty", constant))
    meta = {
        "physical_row_count": physical,
        "data_row_count": data_rows,
        "column_count": width,
        "empty_row_count": empty_rows,
        "empty_columns": empty_columns,
    }
    return headers, meta


class CsvLoader(_DelimitedLoader):
    def __init__(self):
        super().__init__("csv", ",")


class TsvLoader(_DelimitedLoader):
    def __init__(self):
        super().__init__("tsv", "\t")


class TextLoader(_DelimitedLoader):
    def __init__(self):
        super().__init__("txt", None)

    def profile(self, path: Path, max_bytes: int) -> dict:
        profiled = super().profile(path, max_bytes)
        table = profiled["structure"]["tabular"]
        if table["column_count"] <= 1:
            return {
                "warnings": profiled["warnings"],
                "structure": {
                    "text": {
                        "encoding": table["encoding"],
                        "line_count": table["physical_row_count"],
                        "delimiter": None,
                    }
                },
            }
        return profiled


class JsonLoader:
    format = "json"

    def profile(self, path: Path, max_bytes: int) -> dict:
        if path.stat().st_size > max_bytes:
            raise StructureError("file_too_large")
        try:
            text = path.read_text(encoding="utf-8-sig")
            document = json.loads(text)
        except UnicodeError as exc:
            raise StructureError("encoding_unknown") from exc
        except json.JSONDecodeError as exc:
            raise StructureError("invalid_json") from exc
        stats = {
            "depth": 0,
            "array_count": 0,
            "object_count": 0,
            "null_count": 0,
            "primitive_counts": {"number": 0, "text": 0, "boolean": 0, "null": 0},
            "truncated": False,
            "nodes": 0,
        }
        _walk(document, 1, stats)
        keys = list(document)[:200] if isinstance(document, dict) else []
        record_keys = _record_keys(document)
        record_path = None
        if record_keys is None and isinstance(document, dict):
            for key, value in document.items():
                found = _record_keys(value)
                if found is not None:
                    record_keys = found
                    record_path = key
                    break
        root = "object" if isinstance(document, dict) else "array" if isinstance(document, list) else "other"
        return {
            "warnings": ["truncated"] if stats["truncated"] else [],
            "structure": {
                "json": {
                    "root_type": root,
                    "depth": stats["depth"],
                    "keys": keys,
                    "array_count": stats["array_count"],
                    "object_count": stats["object_count"],
                    "null_count": stats["null_count"],
                    "primitive_counts": stats["primitive_counts"],
                    "possible_tabular_records": record_keys is not None,
                    "record_keys": record_keys,
                    "record_path": record_path,
                }
            },
        }


def _walk(node, depth: int, stats: dict) -> None:
    if stats["nodes"] > 100000:
        stats["truncated"] = True
        return
    stats["nodes"] += 1
    stats["depth"] = max(stats["depth"], depth)
    if node is None:
        stats["null_count"] += 1
        stats["primitive_counts"]["null"] += 1
    elif isinstance(node, bool):
        stats["primitive_counts"]["boolean"] += 1
    elif isinstance(node, (int, float)):
        stats["primitive_counts"]["number"] += 1
    elif isinstance(node, str):
        stats["primitive_counts"]["text"] += 1
    elif isinstance(node, list):
        stats["array_count"] += 1
        for item in node:
            _walk(item, depth + 1, stats)
    elif isinstance(node, dict):
        stats["object_count"] += 1
        for value in node.values():
            _walk(value, depth + 1, stats)


def _record_keys(node) -> list[str] | None:
    if not isinstance(node, list) or not node or not all(isinstance(item, dict) for item in node):
        return None
    signatures = [tuple(item.keys()) for item in node]
    if len(set(signatures)) != 1:
        return None
    return list(node[0].keys())


def loader_for(fmt: str):
    if fmt == "xlsx":
        return ExcelLoader("xlsx")
    if fmt == "xlsm":
        return ExcelLoader("xlsm")
    if fmt == "csv":
        return CsvLoader()
    if fmt == "tsv":
        return TsvLoader()
    if fmt == "json":
        return JsonLoader()
    if fmt == "txt":
        return TextLoader()
    raise StructureError("unsupported_extension")


def detect_format(path: Path, original_name: str) -> str:
    suffix = Path(original_name).suffix.lower()
    from app.research.contract import ALLOWED_EXTENSIONS

    if suffix not in ALLOWED_EXTENSIONS:
        raise StructureError("unsupported_extension")
    fmt = ALLOWED_EXTENSIONS[suffix]
    with path.open("rb") as handle:
        head = handle.read(8)
    if fmt in {"xlsx", "xlsm"}:
        if not head.startswith(b"PK\x03\x04"):
            raise StructureError("content_mismatch")
        names = [item.filename for item in _zip_members(path)]
        if not any(name.endswith("workbook.xml") for name in names):
            raise StructureError("content_mismatch")
        return fmt
    if b"\x00" in head:
        raise StructureError("content_mismatch")
    if head.startswith(b"PK\x03\x04"):
        raise StructureError("content_mismatch")
    if fmt == "json":
        try:
            text = path.read_text(encoding="utf-8-sig")
            parsed = json.loads(text)
        except (UnicodeError, json.JSONDecodeError) as exc:
            raise StructureError("invalid_json") from exc
        if not isinstance(parsed, (dict, list)):
            return fmt
        return fmt
    return fmt


def profile_path(path: Path, fmt: str, max_bytes: int) -> dict:
    return loader_for(fmt).profile(path, max_bytes)
