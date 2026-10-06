"""Perfil exploratorio de libros Excel. No limpia ni reinterpreta celdas."""

from __future__ import annotations

import math
import re
import statistics
from datetime import date, datetime
from zipfile import BadZipFile

from openpyxl.utils.exceptions import InvalidFileException

from app.research.schema import TOOL_VERSION

UNIT_PATTERN = re.compile(r"\(([^)]+)\)\s*$")


class WorkbookError(Exception):
    def __init__(self, filename: str):
        self.filename = filename
        super().__init__(filename)


def cell_kind(value) -> str:
    if value is None:
        return "empty"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return "empty"
        return "number"
    if isinstance(value, (datetime, date)):
        return "datetime"
    if isinstance(value, str):
        return "empty" if value.strip() == "" else "text"
    return "other"


def json_value(value):
    kind = cell_kind(value)
    if kind == "empty":
        return None
    if kind == "datetime":
        return value.isoformat()
    if kind in {"number", "boolean"}:
        return value
    if kind == "text":
        return value.strip()
    return str(value)


def classify_row(values) -> str:
    kinds = [cell_kind(value) for value in values]
    texts = sum(kind == "text" for kind in kinds)
    numbers = sum(kind in {"number", "datetime"} for kind in kinds)
    others = sum(kind in {"boolean", "other"} for kind in kinds)
    if texts == 0 and numbers == 0 and others == 0:
        return "empty"
    if texts and numbers <= texts and others == 0:
        return "header"
    return "data"


def used_width(rows) -> int:
    width = 0
    for row in rows:
        for index, value in enumerate(row, start=1):
            if cell_kind(value) != "empty":
                width = max(width, index)
    return width


def read_workbook(path) -> list[dict]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, OSError, KeyError, ValueError) as exc:
        raise WorkbookError(path.name) from exc
    sheets = []
    try:
        for name in workbook.sheetnames:
            worksheet = workbook[name]
            rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
            sheets.append({"name": name, "rows": rows})
    except (InvalidFileException, BadZipFile, OSError, KeyError, ValueError) as exc:
        raise WorkbookError(path.name) from exc
    finally:
        workbook.close()
    return sheets


def _unit(headers) -> str | None:
    for value in reversed(headers):
        if isinstance(value, str):
            match = UNIT_PATTERN.search(value.strip())
            if match:
                return match.group(1)
    return None


def _display_name(headers):
    for value in reversed(headers):
        if isinstance(value, str) and value.strip():
            return value.strip()
    return None


def _duplicate_headers(header_rows) -> list[dict]:
    found = []
    for row_number, row in enumerate(header_rows, start=1):
        counts: dict[str, int] = {}
        for value in row:
            if isinstance(value, str) and value.strip():
                text = value.strip()
                counts[text] = counts.get(text, 0) + 1
        for text, count in counts.items():
            if count > 1:
                found.append({"header_row": row_number, "text": text, "count": count})
    return found


def _column_profile(index: int, headers, data_values) -> dict:
    present = [value for value in data_values if cell_kind(value) != "empty"]
    kinds = {cell_kind(value) for value in present}
    if not present:
        inferred = "empty"
    elif len(kinds) == 1:
        inferred = next(iter(kinds))
    else:
        inferred = "mixed"
    encoded = [json_value(value) for value in present]
    profile = {
        "index": index,
        "headers": [json_value(value) for value in headers],
        "name": _display_name(headers),
        "unit": _unit(headers),
        "count": len(data_values),
        "non_null": len(present),
        "null_count": len(data_values) - len(present),
        "inferred_type": inferred,
        "unique_count": len(set(map(lambda item: repr(item), encoded))),
        "empty": len(present) == 0,
    }
    if encoded:
        profile["first"] = encoded[0]
        profile["last"] = encoded[-1]
    if inferred == "number" and present:
        numbers = [float(value) for value in present]
        profile["min"] = min(numbers)
        profile["max"] = max(numbers)
        profile["mean"] = statistics.mean(numbers)
        profile["median"] = statistics.median(numbers)
    return profile


def _pad(row, width: int):
    clipped = tuple(row[:width])
    return clipped + (None,) * (width - len(clipped))


def _split_blocks(rows) -> tuple[list[dict], int]:
    """Separa regiones por filas de encabezado repetidas. No rellena celdas."""
    blocks = []
    current = None
    empty_rows = 0
    for row_number, row in enumerate(rows, start=1):
        kind = classify_row(row)
        if kind == "empty":
            empty_rows += 1
            continue
        if kind == "header" and (current is None or current["data_rows"]):
            current = {
                "start_row": row_number,
                "end_row": row_number,
                "header_rows": [],
                "data_rows": [],
            }
            blocks.append(current)
        if current is None:
            current = {
                "start_row": row_number,
                "end_row": row_number,
                "header_rows": [],
                "data_rows": [],
            }
            blocks.append(current)
        current["end_row"] = row_number
        if kind == "header":
            current["header_rows"].append(row)
        else:
            current["data_rows"].append(row)
    return blocks, empty_rows


def _profile_block(block: dict) -> dict:
    raw_rows = block["header_rows"] + block["data_rows"]
    width = used_width(raw_rows)
    header_rows = [_pad(row, width) for row in block["header_rows"]]
    data_rows = [_pad(row, width) for row in block["data_rows"]]
    columns = []
    for index in range(width):
        headers = [row[index] for row in header_rows]
        values = [row[index] for row in data_rows]
        columns.append(_column_profile(index + 1, headers, values))
    names = [column["name"] for column in columns if column["name"]]
    name_counts: dict[str, int] = {}
    for item in names:
        name_counts[item] = name_counts.get(item, 0) + 1
    return {
        "start_row": block["start_row"],
        "end_row": block["end_row"],
        "header_row_count": len(header_rows),
        "data_row_count": len(data_rows),
        "column_count": width,
        "duplicate_headers": _duplicate_headers(header_rows),
        "duplicate_column_names": sorted(name for name, count in name_counts.items() if count > 1),
        "empty_columns": [column["index"] for column in columns if column["empty"]],
        "columns": columns,
    }


def profile_sheet(name: str, rows) -> dict:
    blocks, empty_rows = _split_blocks(rows)
    profiled = [_profile_block(block) for block in blocks]
    return {
        "name": name,
        "row_count": len(rows),
        "column_count": used_width(rows),
        "block_count": len(profiled),
        "data_row_count": sum(block["data_row_count"] for block in profiled),
        "empty_row_count": empty_rows,
        "blocks": profiled,
    }


def profile_workbook(path) -> dict:
    sheets = [profile_sheet(item["name"], item["rows"]) for item in read_workbook(path)]
    return {
        "filename": path.name,
        "tool_version": TOOL_VERSION,
        "sheets": sheets,
    }


def preview_sheet(path, sheet_name: str, limit: int, block_index: int = 1) -> dict:
    sheets = {item["name"]: item["rows"] for item in read_workbook(path)}
    if sheet_name not in sheets:
        return {}
    blocks, _empty = _split_blocks(sheets[sheet_name])
    if block_index < 1 or block_index > len(blocks):
        return {"unknown_block": True, "block_count": len(blocks)}
    profiled = _profile_block(blocks[block_index - 1])
    width = profiled["column_count"]
    rows = [
        [json_value(value) for value in _pad(row, width)]
        for row in blocks[block_index - 1]["data_rows"]
    ]
    shown = rows[:limit]
    return {
        "sheet": sheet_name,
        "block": block_index,
        "block_count": len(blocks),
        "limit": limit,
        "shown": len(shown),
        "data_row_count": profiled["data_row_count"],
        "columns": profiled["columns"],
        "rows": shown,
    }
